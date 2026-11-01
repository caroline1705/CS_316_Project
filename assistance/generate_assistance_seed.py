from pathlib import Path
from openpyxl import load_workbook

folder = Path(__file__).resolve().parent
excel_file = folder / "assistance_programs.xlsx"
workbook = load_workbook(excel_file, data_only = True)
sheet = workbook["Coding Ready"]
program_count = sheet.max_row - 1

print(f"Found {program_count} assistance programs.")

headers = [cell.value for cell in sheet[1]]
for index, header in enumerate(headers, start = 1):
    print(f"{index}. {header}")

programs = []
for row in sheet.iter_rows(min_row=2, values_only=True):
    program = dict(zip(headers, row))
    programs.append(program)

print("First program ID:", programs[0]["program_id"])
print("First program name:", programs[0]["program_name"])
print("First program FPL limit:", programs[0]["income_limit_fpl"])

main_columns = [
    "program_id",
    "program_name",
    "organization",
    "residency_requirements",
    "additional_insurance_requirements",
    "income_limit_fpl",
    "other_income_requirements",
    "assistance_type",
    "assistance_amount",
    "other_eligibility_requirements",
    "application_url",
    "source_url",
    "notes",
    "insurance_coverage",
    "min_age",
    "max_age"
]
for column in main_columns:
    if column not in headers:
        raise ValueError(f"Missing spreadsheet column: {column}")

print("All 16 main table columns were found.")
    
def sql_value(value):
    if value is None:
        return "NULL"
    if isinstance(value, str):
        escaped = value.replace("'", "''")
        return f"'{escaped}'"
    if isinstance(value, (int, float)):
        return str(value)
    raise ValueError(f"Unsupported value type: {type(value)}")

##print("Text test:", sql_value("Pink Fund"))
##print("Number test:", sql_value(300))
##print("Blank test:", sql_value(None))
##print("Apostrophe test:", sql_value("year's income"))

insert_statements = []
for program in programs:
    values = [sql_value(program[column]) for column in main_columns]
    column_names = ", ".join(main_columns)
    sql_values = ", ".join(values)
    insert_sql = (
        f"INSERT INTO assistance_program ({column_names})\n"
        f"VALUES ({sql_values});"
    )
    insert_statements.append(insert_sql)

print(f"\nGenerated {len(insert_statements)} INSERT statements.")
print("\nFirst program:")
print(insert_statements[0])
print("\nLast program:")
print(insert_statements[-1])

program_ids = [program["program_id"] for program in programs]

unique_ids = set(program_ids)

print("\nTotal program IDs:", len(program_ids))
print("Unique program IDs:", len(unique_ids))

if len(program_ids) != len(unique_ids):
    raise ValueError("Duplicate program IDs found!")

print("All program IDs are unique.")

##Work for cancer table

print("\nSample cancer and insurance values:")

for program in programs[:5]:
    print(
        program["program_id"],
        "| Cancer:", repr(program["cancer_types"]),
        "| Subtype:", repr(program["cancer_subtype"]),
        "| Insurance:", repr(program["accepted_insurance_types"])
    )

print("\nAll cancer and insurance values:")

for program in programs:
    print(
        program["program_id"],
        "| Cancer:", repr(program["cancer_types"]),
        "| Subtype:", repr(program["cancer_subtype"]),
        "| Insurance:", repr(program["accepted_insurance_types"])
    )

def split_values(value):
    if value is None:
        return []
    return [
        item.strip()
        for item in value.split(",")
        if item.strip()
    ]

print("\nInsurance splitting test:")
print(split_values("Medicare, Medicaid, Private"))

print("Blank insurance test:")
print(split_values(None))

print("Single cancer type test:")
print(split_values("Breast"))

cancer_insert_statements = []

for program in programs:
    program_id = program["program_id"]
    cancer_types = split_values(program["cancer_types"])
    cancer_subtype = program["cancer_subtype"] or ""

    for cancer_type in cancer_types:
        insert_sql = (
            "INSERT INTO assistance_program_cancer "
            "(program_id, cancer_type, cancer_subtype)\n"
            f"VALUES ({sql_value(program_id)}, "
            f"{sql_value(cancer_type)}, "
            f"{sql_value(cancer_subtype)});"
        )

        cancer_insert_statements.append(insert_sql)

print(f"\nGenerated {len(cancer_insert_statements)} cancer INSERT statements.")
print("First cancer INSERT:")
print(cancer_insert_statements[0])
print("Last cancer INSERT:")
print(cancer_insert_statements[-1])

##Work for insurance table

insurance_insert_statements = []
for program in programs:
    program_id = program["program_id"]
    insurance_types = split_values(program["accepted_insurance_types"])
    for insurance_type in insurance_types:
        insert_sql = (
            "INSERT INTO assistance_program_insurance "
            "(program_id, insurance_type)\n"
            f"VALUES ({sql_value(program_id)}, "
            f"{sql_value(insurance_type)});"
        )
        insurance_insert_statements.append(insert_sql)
print(f"\nGenerated {len(insurance_insert_statements)} insurance INSERT statements.")
print("First insurance INSERT:")
print(insurance_insert_statements[0])
print("Last insurance INSERT:")
print(insurance_insert_statements[-1])

all_statements = (
    insert_statements
    + cancer_insert_statements
    + insurance_insert_statements
)
output_file = folder / "generated_assistance_seed.sql"
sql_content = "\n\n".join(all_statements) + "\n"
output_file.write_text(sql_content, encoding="utf-8")
print(f"\nSaved {len(all_statements)} INSERT statements to {output_file.name}")